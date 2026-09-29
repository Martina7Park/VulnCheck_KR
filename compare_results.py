#!/usr/bin/env python3
"""
점검 결과 diff / 추세 비교 도구

[사용법]
  1) 두 엑셀 비교 (이전 vs 현재):
     python compare_results.py diff  이전결과.xlsx  현재결과.xlsx

  2) 감사 로그 기반 추세 분석:
     python compare_results.py trend  logs/

  3) 특정 호스트만 비교:
     python compare_results.py diff  이전.xlsx  현재.xlsx  --host WEB-01

[출력]
  - 콘솔 요약
  - diff → compare_diff_YYYYMMDD_HHMMSS.xlsx
  - trend → compare_trend_YYYYMMDD_HHMMSS.xlsx
"""

import os, sys, datetime, json, glob, re

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except ImportError:
    print("[ERROR] openpyxl 필요: pip install openpyxl")
    sys.exit(1)


RESULTS = {"양호", "취약", "수동확인", "N-A"}
FONT_NAME = "맑은 고딕"

FILL_GOOD   = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_BAD    = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FILL_NA     = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
FILL_MANUAL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FILL_NEW_VULN = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")
FILL_RESOLVED = PatternFill(start_color="44BB44", end_color="44BB44", fill_type="solid")
FILL_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

RESULT_FILL = {
    "양호": FILL_GOOD, "취약": FILL_BAD,
    "N-A": FILL_NA, "수동확인": FILL_MANUAL,
}

BD = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def read_text(path):
    for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr", "latin-1"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    return open(path, encoding="latin-1").read()


# ================================================================
# 엑셀 파싱: {(hostname, item_code): {"result": str, "reason": str}}
# ================================================================
def parse_result_excel(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    header_row = None
    id_col = result_col = reason_col = host_col = None

    for row_idx in range(1, min(20, ws.max_row + 1)):
        for col_idx in range(1, ws.max_column + 1):
            v = ws.cell(row_idx, col_idx).value
            if not v:
                continue
            sv = str(v).strip().lower()
            if "평가항목id" in sv or "항목id" in sv or "항목코드" in sv:
                header_row = row_idx
                id_col = col_idx
                break
        if header_row:
            break

    if not header_row:
        print(f"  [WARN] 항목ID 헤더를 찾을 수 없음: {xlsx_path}")
        wb.close()
        return {}

    for col_idx in range(1, ws.max_column + 1):
        v = ws.cell(header_row, col_idx).value
        if not v:
            continue
        sv = str(v).strip().lower()
        if ("점검결과" in sv or "결과" in sv) and result_col is None:
            result_col = col_idx
        if ("현황" in sv or "근거" in sv) and reason_col is None:
            reason_col = col_idx
        if "호스트" in sv or "장비" in sv or "서버" in sv or "대상" in sv:
            host_col = col_idx

    data = {}
    for row_idx in range(header_row + 1, ws.max_row + 1):
        code_val = ws.cell(row_idx, id_col).value
        if not code_val:
            continue
        code = str(code_val).strip().upper()

        result_val = ""
        if result_col:
            rv = ws.cell(row_idx, result_col).value
            if rv:
                result_val = str(rv).strip()

        reason_val = ""
        if reason_col:
            rv = ws.cell(row_idx, reason_col).value
            if rv:
                reason_val = str(rv).strip()

        hostname = "UNKNOWN"
        if host_col:
            hv = ws.cell(row_idx, host_col).value
            if hv:
                hostname = str(hv).strip()
        if hostname == "UNKNOWN" and reason_val:
            m = re.match(r"\[([^\]]+)\]", reason_val)
            if m:
                hostname = m.group(1)

        data[(hostname, code)] = {"result": result_val, "reason": reason_val}

    wb.close()
    return data


# ================================================================
# diff 모드: 이전 vs 현재 비교
# ================================================================
def run_diff(prev_path, curr_path, host_filter=None):
    print(f"[비교] 이전: {os.path.basename(prev_path)}")
    print(f"[비교] 현재: {os.path.basename(curr_path)}")

    prev = parse_result_excel(prev_path)
    curr = parse_result_excel(curr_path)

    if host_filter:
        prev = {k: v for k, v in prev.items() if k[0] == host_filter}
        curr = {k: v for k, v in curr.items() if k[0] == host_filter}

    all_keys = sorted(set(prev.keys()) | set(curr.keys()))

    new_vulns = []
    resolved = []
    changed = []
    unchanged = []
    new_items = []
    removed_items = []

    for key in all_keys:
        hostname, code = key
        p = prev.get(key)
        c = curr.get(key)

        if p and not c:
            removed_items.append({"host": hostname, "code": code,
                                  "prev_result": p["result"], "prev_reason": p["reason"]})
            continue
        if not p and c:
            new_items.append({"host": hostname, "code": code,
                              "curr_result": c["result"], "curr_reason": c["reason"]})
            if c["result"] == "취약":
                new_vulns.append({"host": hostname, "code": code,
                                  "reason": c["reason"]})
            continue

        pr = p["result"]
        cr = c["result"]

        if pr == cr:
            unchanged.append({"host": hostname, "code": code, "result": cr})
        else:
            entry = {"host": hostname, "code": code,
                     "prev": pr, "curr": cr,
                     "prev_reason": p["reason"], "curr_reason": c["reason"]}
            changed.append(entry)
            if pr != "취약" and cr == "취약":
                new_vulns.append({"host": hostname, "code": code,
                                  "reason": c["reason"]})
            if pr == "취약" and cr != "취약":
                resolved.append({"host": hostname, "code": code,
                                 "new_result": cr, "reason": c["reason"]})

    # 콘솔 요약
    print(f"\n{'='*60}")
    print(f" 비교 결과 요약")
    print(f"{'='*60}")
    print(f"  전체 항목:       {len(all_keys):>5}건")
    print(f"  변경 없음:       {len(unchanged):>5}건")
    print(f"  결과 변경:       {len(changed):>5}건")
    print(f"  신규 취약:       {len(new_vulns):>5}건  ← 주의")
    print(f"  취약→조치완료:   {len(resolved):>5}건  ← 개선")
    print(f"  신규 추가 항목:  {len(new_items):>5}건")
    print(f"  삭제된 항목:     {len(removed_items):>5}건")
    print(f"{'='*60}")

    if new_vulns:
        print(f"\n▼ 신규 취약 항목 ({len(new_vulns)}건):")
        for v in new_vulns:
            print(f"  [{v['host']}] {v['code']} — {v['reason'][:60]}")

    if resolved:
        print(f"\n▲ 조치 완료 항목 ({len(resolved)}건):")
        for v in resolved:
            print(f"  [{v['host']}] {v['code']} → {v['new_result']}")

    # 엑셀 출력
    wb = openpyxl.Workbook()

    # 시트 1: 요약
    ws_sum = wb.active
    ws_sum.title = "비교 요약"
    _write_summary_sheet(ws_sum, prev_path, curr_path,
                         len(all_keys), len(unchanged), len(changed),
                         len(new_vulns), len(resolved),
                         len(new_items), len(removed_items))

    # 시트 2: 신규 취약
    if new_vulns:
        ws_nv = wb.create_sheet("신규 취약")
        _write_list_sheet(ws_nv, ["호스트", "항목코드", "근거"],
                          [[v["host"], v["code"], v["reason"]] for v in new_vulns],
                          FILL_NEW_VULN)

    # 시트 3: 조치 완료
    if resolved:
        ws_rv = wb.create_sheet("조치 완료")
        _write_list_sheet(ws_rv, ["호스트", "항목코드", "변경 후 결과", "근거"],
                          [[v["host"], v["code"], v["new_result"], v["reason"]] for v in resolved],
                          FILL_RESOLVED)

    # 시트 4: 전체 변경
    if changed:
        ws_ch = wb.create_sheet("결과 변경")
        _write_list_sheet(ws_ch, ["호스트", "항목코드", "이전 결과", "현재 결과", "이전 근거", "현재 근거"],
                          [[e["host"], e["code"], e["prev"], e["curr"],
                            e["prev_reason"], e["curr_reason"]] for e in changed])

    # 시트 5: 전체 diff (호스트×항목 매트릭스)
    ws_all = wb.create_sheet("전체 비교")
    _write_full_diff_sheet(ws_all, all_keys, prev, curr)

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = os.path.dirname(curr_path) or "."
    out_path = os.path.join(out_dir, f"compare_diff_{ts}.xlsx")
    wb.save(out_path)
    print(f"\n[완료] 비교 결과 저장: {out_path}")
    return out_path


def _write_summary_sheet(ws, prev_name, curr_name,
                         total, unchanged, chg, new_v, resolved, new_i, removed):
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 50

    rows = [
        ("비교 일시", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("이전 결과", os.path.basename(prev_name)),
        ("현재 결과", os.path.basename(curr_name)),
        ("", ""),
        ("전체 항목", total),
        ("변경 없음", unchanged),
        ("결과 변경", chg),
        ("신규 취약", new_v),
        ("취약→조치완료", resolved),
        ("신규 추가 항목", new_i),
        ("삭제된 항목", removed),
    ]
    for i, (label, val) in enumerate(rows, 1):
        ws.cell(i, 1, label).font = Font(name=FONT_NAME, size=10, bold=True)
        ws.cell(i, 2, val).font = Font(name=FONT_NAME, size=10)
        if label == "신규 취약" and val > 0:
            ws.cell(i, 2).fill = FILL_BAD
        if label == "취약→조치완료" and val > 0:
            ws.cell(i, 2).fill = FILL_GOOD


def _write_list_sheet(ws, headers, rows, header_fill=None):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        c.fill = header_fill or FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BD

    for ri, row in enumerate(rows, 2):
        for ci, val in enumerate(row, 1):
            c = ws.cell(ri, ci, val)
            c.font = Font(name=FONT_NAME, size=9)
            c.border = BD
            c.alignment = Alignment(vertical="center", wrap_text=True)

    for ci in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = max(15, len(headers[ci-1]) * 2 + 4)
    if len(headers) >= 3:
        ws.column_dimensions[openpyxl.utils.get_column_letter(len(headers))].width = 60


def _write_full_diff_sheet(ws, all_keys, prev, curr):
    headers = ["호스트", "항목코드", "이전 결과", "현재 결과", "변경", "현재 근거"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")
        c.border = BD

    ri = 2
    for key in sorted(all_keys):
        hostname, code = key
        p = prev.get(key, {})
        c_data = curr.get(key, {})
        pr = p.get("result", "-")
        cr = c_data.get("result", "-")

        if pr == cr:
            change_label = ""
        elif pr == "-":
            change_label = "신규"
        elif cr == "-":
            change_label = "삭제"
        elif pr != "취약" and cr == "취약":
            change_label = "악화"
        elif pr == "취약" and cr != "취약":
            change_label = "개선"
        else:
            change_label = "변경"

        row_vals = [hostname, code, pr, cr, change_label, c_data.get("reason", "")]
        for ci, val in enumerate(row_vals, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = Font(name=FONT_NAME, size=9)
            cell.border = BD
            if ci == 3 and val in RESULT_FILL:
                cell.fill = RESULT_FILL[val]
            if ci == 4 and val in RESULT_FILL:
                cell.fill = RESULT_FILL[val]
            if ci == 5:
                if val == "악화":
                    cell.fill = FILL_BAD
                    cell.font = Font(name=FONT_NAME, size=9, bold=True, color="9C0006")
                elif val == "개선":
                    cell.fill = FILL_GOOD
                    cell.font = Font(name=FONT_NAME, size=9, bold=True, color="006100")
        ri += 1

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 60
    ws.auto_filter.ref = f"A1:F{ri - 1}"


# ================================================================
# trend 모드: 감사 로그 기반 추세 분석
# ================================================================
def run_trend(log_dir):
    print(f"[추세] 로그 폴더: {log_dir}")

    log_files = sorted(glob.glob(os.path.join(log_dir, "*.json")))
    if not log_files:
        print("[ERROR] JSON 로그 파일 없음")
        return

    entries = []
    for lf in log_files:
        try:
            text = read_text(lf)
            data = json.loads(text)
            entries.append(data)
        except (json.JSONDecodeError, Exception) as e:
            print(f"  [WARN] 로그 파싱 실패: {lf} — {e}")

    if not entries:
        print("[ERROR] 유효한 로그 없음")
        return

    entries.sort(key=lambda e: e.get("timestamp", ""))

    print(f"\n{'='*60}")
    print(f" 추세 분석 ({len(entries)}건 로그)")
    print(f"{'='*60}")
    print(f"  {'일시':<22} {'기준':<8} {'카테고리':<8} {'대상':>4} {'양호':>5} {'취약':>5} {'N-A':>5} {'수동':>5} {'취약률':>7}")
    print(f"  {'-'*22} {'-'*8} {'-'*8} {'-'*4} {'-'*5} {'-'*5} {'-'*5} {'-'*5} {'-'*7}")

    trend_data = []
    for e in entries:
        s = e.get("summary", {})
        ts = e.get("timestamp", "?")[:19]
        mode = e.get("mode_code", "?")
        cat = e.get("category", "?")
        tc = e.get("target_count", 0)
        good = s.get("good", 0)
        bad = s.get("bad", 0)
        na = s.get("na", 0)
        mc = s.get("manual", 0)
        total = good + bad + na + mc
        rate = f"{bad/total*100:.1f}%" if total > 0 else "-"

        print(f"  {ts:<22} {mode:<8} {cat:<8} {tc:>4} {good:>5} {bad:>5} {na:>5} {mc:>5} {rate:>7}")
        trend_data.append({
            "timestamp": ts, "mode": mode, "category": cat,
            "targets": tc, "good": good, "bad": bad, "na": na, "manual": mc,
            "total": total, "rate": rate,
        })

    # 카테고리별 추세
    cats = {}
    for td in trend_data:
        key = f"{td['mode']}_{td['category']}"
        cats.setdefault(key, []).append(td)

    print(f"\n{'='*60}")
    print(f" 카테고리별 추세 변화")
    print(f"{'='*60}")
    for key, items in cats.items():
        if len(items) < 2:
            continue
        first = items[0]
        last = items[-1]
        delta_bad = last["bad"] - first["bad"]
        direction = "↑악화" if delta_bad > 0 else "↓개선" if delta_bad < 0 else "→유지"
        print(f"  {key}: 취약 {first['bad']}건 → {last['bad']}건 ({direction}, Δ{delta_bad:+d})")

    # 엑셀 출력
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "추세 데이터"

    headers = ["일시", "기준", "카테고리", "대상 수", "양호", "취약", "N-A", "수동확인", "전체", "취약률"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")
        c.border = BD

    for ri, td in enumerate(trend_data, 2):
        vals = [td["timestamp"], td["mode"], td["category"], td["targets"],
                td["good"], td["bad"], td["na"], td["manual"], td["total"], td["rate"]]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(ri, ci, val)
            cell.font = Font(name=FONT_NAME, size=9)
            cell.border = BD
            if ci == 6 and isinstance(val, int) and val > 0:
                cell.fill = FILL_BAD

    for ci, w in enumerate([22, 10, 10, 8, 8, 8, 8, 8, 8, 10], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    # 카테고리별 요약 시트
    if cats:
        ws2 = wb.create_sheet("카테고리 추세")
        headers2 = ["카테고리", "첫 점검일", "마지막 점검일", "점검 횟수",
                     "첫 취약", "현재 취약", "변화", "방향"]
        for ci, h in enumerate(headers2, 1):
            c = ws2.cell(1, ci, h)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            c.fill = FILL_HEADER
            c.alignment = Alignment(horizontal="center")
            c.border = BD

        ri2 = 2
        for key, items in cats.items():
            first = items[0]
            last = items[-1]
            delta = last["bad"] - first["bad"]
            direction = "악화" if delta > 0 else "개선" if delta < 0 else "유지"
            vals = [key, first["timestamp"], last["timestamp"], len(items),
                    first["bad"], last["bad"], delta, direction]
            for ci, val in enumerate(vals, 1):
                cell = ws2.cell(ri2, ci, val)
                cell.font = Font(name=FONT_NAME, size=9)
                cell.border = BD
                if ci == 8:
                    if val == "악화":
                        cell.fill = FILL_BAD
                    elif val == "개선":
                        cell.fill = FILL_GOOD
            ri2 += 1

        for ci, w in enumerate([18, 22, 22, 10, 10, 10, 8, 8], 1):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(os.path.dirname(log_dir.rstrip("/\\")), f"compare_trend_{ts}.xlsx")
    wb.save(out_path)
    print(f"\n[완료] 추세 결과 저장: {out_path}")
    return out_path


# ================================================================
# CLI
# ================================================================
def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(0)

    cmd = sys.argv[1].lower()

    if cmd == "diff":
        if len(sys.argv) < 4:
            print("사용법: python compare_results.py diff <이전.xlsx> <현재.xlsx> [--host HOSTNAME]")
            sys.exit(1)
        prev_path = sys.argv[2]
        curr_path = sys.argv[3]
        host_filter = None
        if "--host" in sys.argv:
            idx = sys.argv.index("--host")
            if idx + 1 < len(sys.argv):
                host_filter = sys.argv[idx + 1]
        run_diff(prev_path, curr_path, host_filter)

    elif cmd == "trend":
        if len(sys.argv) < 3:
            print("사용법: python compare_results.py trend <logs_폴더>")
            sys.exit(1)
        log_dir = sys.argv[2]
        run_trend(log_dir)

    else:
        print(f"[ERROR] 알 수 없는 명령: {cmd}")
        print("  사용 가능: diff, trend")
        sys.exit(1)


if __name__ == "__main__":
    main()
